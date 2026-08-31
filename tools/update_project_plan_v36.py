"""Create the V3.6 plan update for distributed FS-IR02 liquid-level stations."""
from pathlib import Path

from docx import Document

SOURCE = Path(r"D:\\shixi\\docs\\综合管廊数字孪生运维实体样品项目计划书_V3.5_硬件数量与型号冻结版.docx")
OUTPUT = Path(r"D:\\shixi\\docs\\综合管廊数字孪生运维实体样品项目计划书_V3.6_液位监测与水路建模版.docx")


def set_cell(cell, value):
    cell.text = value


def row_by_code(table, code):
    for row in table.rows[1:]:
        if row.cells[0].text.strip() == code:
            return row
    raise KeyError(code)


def set_bom(table, code, name, qty, purpose, estimate):
    row = row_by_code(table, code)
    for cell, value in zip(row.cells, (code, name, qty, purpose, estimate)):
        set_cell(cell, value)


def add_bom(table, code, name, qty, purpose, estimate):
    row = table.add_row()
    for cell, value in zip(row.cells, (code, name, qty, purpose, estimate)):
        set_cell(cell, value)


def add_interface(table, module, pin, electrical, constraint):
    row = table.add_row()
    for cell, value in zip(row.cells, (module, pin, electrical, constraint)):
        set_cell(cell, value)


def add_status(table, domain, status, fact, gap):
    row = table.add_row()
    for cell, value in zip(row.cells, (domain, status, fact, gap)):
        set_cell(cell, value)


def add_ledger(table, asset, hardware, purpose, limits):
    row = table.add_row()
    for cell, value in zip(row.cells, (asset, hardware, purpose, limits)):
        set_cell(cell, value)


def set_ledger(table, asset, hardware, purpose, limits):
    row = row_by_code(table, asset)
    for cell, value in zip(row.cells, (asset, hardware, purpose, limits)):
        set_cell(cell, value)


doc = Document(SOURCE)

# Keep the established layout and make only version-local edits.
for section in doc.sections:
    for paragraph in section.header.paragraphs:
        for run in paragraph.runs:
            if "项目计划书 V3.5" in run.text:
                run.text = run.text.replace("项目计划书 V3.5", "项目计划书 V3.6")

table0 = doc.tables[0]
set_cell(table0.rows[1].cells[1], "V3.6（液位监测与水路建模版）")
set_cell(
    table0.rows[6].cells[1],
    "V3.6新增 FS-IR02 液位传感器×5，并固化 V17 模型中的分散安装、120 mm 风机和水路设计。"
    "实体接线、孔径和泵运行仍以实物核验与安全评审为准。",
)

# BOM: retain the five units as distinct hardware kits and state what is known.
bom = doc.tables[21]
add_bom(
    bom, "H-20A", "FS-IR02 光学液位传感器套件（控制板+外置探头）", "5",
    "分布式低液位、常规液位、高液位、超高液位和泄漏位监测；板卡约38.6×22.1mm，5V，XH2.54。",
    "待报价",
)
set_bom(
    bom, "H-12", "12V四线轴流风机（120×120mm）", "2",
    "进气与排风各1台；确认PWM、TACH、额定电流、转向和120mm安装孔距，MCU不得直接驱动。", "待报价",
)
set_bom(
    bom, "H-18", "JQC-3FF-S-Z 继电器模块（1路）", "1",
    "V3.6设计为24V水泵P-01低压通断；先核验线圈电压、触发电平、触点额定值和感性负载保护。", "待报价",
)
set_bom(
    bom, "H-21A", "24V水泵 P-01", "1",
    "V17模型定义为水槽→P-01→V-01→补水歧管→溢流回水槽的低水量液位测试回路；未投运。", "待报价",
)
set_bom(
    bom, "H-32", "24V水泵驱动、保险、浪涌保护与低压水路件", "1套（待实物核验）",
    "仅在完成水电隔离、保险、止回/截止阀、软管接口、漏电/渗漏和干湿分区评审后接入。", "待设计",
)

# Safety and interfaces: explicit 5 V / 3.3 V boundary, no false implementation claim.
logic = doc.tables[22]
set_cell(
    logic.rows[3].cells[1],
    "V3.6已将24V水泵作为低水量液位测试回路的建模对象：水槽→P-01→V-01→补水歧管→溢流回水槽。"
    "该路径尚未实体投运；任何接线前必须完成独立24V保险、水电隔离、止回/截流、泄漏检查和项目负责人批准。"
    "不得连接建筑给水，演示仍限制于独立水槽内少量清水。",
)
interface = doc.tables[23]
add_interface(
    interface, "FS-IR02液位传感器×5", "每套：GND / DO / AO / VCC", "5V；DO与AO待接入多路采集/调理",
    "PCB尺寸约38.6×22.1mm；AO不得直接进入3.3V MCU ADC。探头螺纹、孔径和线长须到货实测后冻结。",
)
add_interface(
    interface, "P-01 24V水泵 / K-01 JQC-3FF-S-Z", "K-01低压控制、P-01独立24V支路", "独立24V、保险、浪涌/续流保护",
    "模型仅表达设计关系；未完成额定电流、触点能力、软管和渗漏测试前不得通水或通电运行。",
)

status = doc.tables[46]
add_status(
    status, "FS-IR02液位监测与V17模型", "建模完成，实体待接入",
    "FS-IR02×5已按L01低液位、L02常规、L03高、L04超高、L05泄漏位分散到V17模型；每站含板卡、探头、支架和线束编号。",
    "待实物测量探头螺纹/孔径、完成5V供电和DO/AO逻辑验证，并设计多通道采集与干湿隔离。",
)
add_status(
    status, "120mm风机与24V水路设计", "模型完成，实体待核验",
    "两台120×120mm四线风机、P-01、V-01、补水歧管与K-01继电器控制关系已在V17建模。",
    "待确认风机孔距/PWM/TACH、泵电流/接口、继电器能力、软管规格和水电安全评审。",
)

priority = doc.tables[53]
set_cell(priority.rows[1].cells[1], "FS-IR02液位传感器×5、传感器端子/线束、五路DO/AO采集与5V/3.3V调理、甲烷/CO/氧气等核心传感器")
set_cell(priority.rows[2].cells[1], "120mm四线风机×2、JQC-3FF-S-Z、24V泵保护、软管、V-01截止阀、补水歧管和干湿隔离件")
set_cell(priority.rows[2].cells[2], "支撑V17分散液位站、通风和低水量水路设计；所有器件仍须数据手册和台架安全验证。")

ledger = doc.tables[62]
set_ledger(
    ledger, "H-08", "JQC-3FF-S-Z 1路继电器模块（K-01，1个）", "P-01 24V水泵低压通断（设计）",
    "继电器触点与MCU低压侧隔离；须核验线圈/触发电平、触点额定值、泵启动电流及感性负载保护后才可接入。",
)
set_ledger(
    ledger, "H-09", "12V四线轴流风机（120×120mm，2个）", "进气/排风联动与风机故障场景",
    "逐台确认额定电压、电流、PWM、TACH、转向和120mm安装孔距；MCU不得直接带电机，须配置驱动与反馈。",
)
set_ledger(
    ledger, "H-30", "24V水泵 P-01（1个）", "V17低水量液位测试回路的执行器（设计，未投运）",
    "水槽→P-01→V-01→补水歧管→溢流回水槽；不得接建筑给水。完成24V保险、水电隔离、渗漏和绝缘检查后才可通水/通电。",
)
add_ledger(
    ledger, "H-33", "FS-IR02光学液位传感器套件（5个）", "L01低、L02常规、L03高、L04超高、L05泄漏位",
    "5V控制板约38.6×22.1mm；DO/AO须调理至MCU电平，探头孔径/螺纹按实物复核，板卡高于最高水位。",
)
add_ledger(
    ledger, "H-34", "V17水路件：P-01、V-01、补水歧管、回流/溢流软管", "独立低水量液位测试回路（设计）",
    "不得连接建筑给水；须完成24V保险、继电器触点、止回/截流、渗漏、绝缘和干湿隔离检查后才可投运。",
)

doc.add_heading("V3.6 液位监测与水路建模变更说明", level=1)
doc.add_paragraph("本版将FS-IR02液位传感器套件新增为5个，并以V17环形模型表达独立安装站和线束。已确认的板卡尺寸为约38.6×22.1mm；探头螺纹、孔径及软管接口没有可靠尺寸来源，必须在实物到位后卡尺复核。")
doc.add_paragraph("V17中水泵循环路径是结构与接线规划，不代表已通水、已接电或已通过安全验收。实体实施必须保持水电隔离、独立低压保险和少量清水演示边界。")

doc.core_properties.title = "综合管廊数字孪生运维实体样品项目计划书 V3.6"
doc.core_properties.subject = "液位监测与水路建模"
doc.core_properties.comments = "V3.6：新增FS-IR02液位传感器×5，更新120mm风机、24V水泵和V17模型设计边界。"
doc.save(OUTPUT)
print(OUTPUT)
