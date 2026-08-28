from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "综合管廊数字孪生运维实体样品项目计划书_V2.7.docx"
OUTPUT = ROOT / "docs" / "综合管廊数字孪生运维实体样品项目计划书_V2.8.docx"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf" if bold else r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def set_paragraph_text(paragraph, value: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def set_cell_text(cell, value: str) -> None:
    set_paragraph_text(cell.paragraphs[0], value)
    for paragraph in cell.paragraphs[1:]:
        set_paragraph_text(paragraph, "")


def replace_everywhere(document: Document, old: str, new: str) -> None:
    paragraphs = list(document.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    for paragraph in paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


def append_styled_row(table, values: tuple[str, str, str]) -> None:
    template = deepcopy(table.rows[-1]._tr)
    table._tbl.append(template)
    row = table.rows[-1]
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)


def make_layout_diagram() -> bytes:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((800, 45), "1 m实体样品冻结布局（俯视示意）", anchor="ma", font=font(38, True), fill="#0B2545")

    outer = (120, 150, 1480, 700)
    draw.rounded_rectangle(outer, radius=20, fill="#F7F8FA", outline="#1F4D78", width=5)
    usable_left, usable_right = 180, 1420
    usable_width = usable_right - usable_left
    zone_lengths = (250, 300, 350)
    zone_names = ("A区 250 mm：入口与环境", "B区 300 mm：渗水监测", "C区 350 mm：气体与设备")
    colors = ("#E8EEF5", "#E2F0D9", "#FFF2CC")
    cursor = usable_left
    boundaries = [cursor]
    for length, name, color in zip(zone_lengths, zone_names, colors):
        zone_width = usable_width * length / sum(zone_lengths)
        end = cursor + zone_width
        draw.rectangle((cursor, outer[1], end, outer[3]), fill=color, outline="#55708D", width=2)
        draw.text(((cursor + end) / 2, 185), name, anchor="ma", font=font(23, True), fill="#1F4D78")
        cursor = end
        boundaries.append(cursor)

    draw.line((180, 310, 1420, 310), fill="#727D8A", width=24)
    draw.text((800, 280), "干燥管廊管线 / B区少量受控滴水点", anchor="ms", font=font(20), fill="#3A3A3A")
    draw.line((180, 470, 1420, 470), fill="#B5651D", width=20)
    draw.text((800, 440), "多气体监测安装区（仅安全信号模拟）", anchor="ms", font=font(20), fill="#7A3E00")
    draw.line((180, 590, 1420, 590), fill="#5A5A5A", width=16)
    draw.text((800, 560), "模拟电缆桥架与分区照明", anchor="ms", font=font(20), fill="#3A3A3A")

    items = [
        (245, 390, "门磁 / 温湿度"),
        (665, 390, "280×180×25 mm可拆接水盘"),
        (1110, 390, "CH₄ / CO / O₂ / 烟雾 / 风机"),
        (1195, 645, "底座下方独立干式电气舱"),
    ]
    for x, y, label in items:
        draw.ellipse((x - 13, y - 13, x + 13, y + 13), fill="#C62828")
        draw.text((x + 22, y), label, anchor="lm", font=font(18, True), fill="#263238")

    draw.text(
        (800, 765),
        "外形 1000 × 500 × 420 mm｜内部净空 900 × 360 × 300 mm｜一体底座 + 可拆4 mm透明罩",
        anchor="mm",
        font=font(22, True),
        fill="#334E68",
    )
    draw.text(
        (800, 815),
        "底座高80 mm；220V适配器置于模型外，模型内部仅使用12V / 5V / 3.3V直流",
        anchor="mm",
        font=font(21),
        fill="#5B6573",
    )

    stream = BytesIO()
    image.save(stream, format="PNG", optimize=True)
    return stream.getvalue()


def main() -> None:
    document = Document(SOURCE)

    replace_everywhere(document, "项目计划书 V2.7", "项目计划书 V2.8")
    replace_everywhere(document, "V2.7（工程二审修订与可验收MVP基线版）", "V2.8（1 m实体尺寸冻结版）")
    replace_everywhere(
        document,
        "工程二审完成；MVP范围、硬件接口、云端边界与里程碑待按本版冻结",
        "工程二审完成；1 m实体外形、内部净空、分区尺寸与水电隔离方案按本版冻结",
    )
    replace_everywhere(
        document,
        "两段式实体可通过普通教室门并由两人搬运；Web平台提供Windows启动脚本和Docker可选方案。",
        "1 m一体式实体可通过普通教室门；建议两人搬运。可拆透明罩、水盘和电气舱便于维护；Web平台提供Windows启动脚本和Docker可选方案。",
    )
    replace_everywhere(document, "两段式结构完成", "一体底座与可拆罩结构完成")
    replace_everywhere(
        document,
        "先简模验证；两段式设计；提供可手工修正孔位和备用支架",
        "先做1:1纸板样机校核；一体底座；按总装±5 mm、孔位±0.5 mm控制并预留1-2 mm装配间隙",
    )

    dimensions_table = next(
        table
        for table in document.tables
        if any(cell.text.strip() == "总尺寸" for row in table.rows for cell in row.cells)
    )
    by_label = {row.cells[0].text.strip(): row for row in dimensions_table.rows[1:]}
    set_cell_text(by_label["总尺寸"].cells[1], "约1000 × 500 × 420 mm")
    set_cell_text(by_label["总尺寸"].cells[2], "桌面展示尺寸；总装外形允许±5 mm，长度以约1 m为冻结基准")
    set_cell_text(by_label["舱体"].cells[1], "单舱剖视；A/B/C三纵向区域")
    set_cell_text(by_label["舱体"].cells[2], "内部有效展示净空约900 × 360 × 300 mm，不构造其他独立系统")
    set_cell_text(by_label["外壳"].cells[1], "4 mm透明亚克力顶罩与观众侧板")
    set_cell_text(by_label["外壳"].cells[2], "分片可拆；板间预留1-2 mm装配间隙，边缘倒钝并限制观众接触水盘和电路")
    set_cell_text(by_label["底座"].cells[1], "1000 × 500 × 80 mm一体式封闭底座")
    set_cell_text(by_label["底座"].cells[2], "B区设可拆高边接水盘；A/C区下方设置独立干式电气舱，水电分隔")
    set_cell_text(by_label["重量"].cells[1], "目标不超过18 kg")
    set_cell_text(by_label["重量"].cells[2], "两人搬运更安全；最终以加工图和材料实重核算")

    append_styled_row(dimensions_table, ("纵向分区", "A 250 / B 300 / C 350 mm", "三段合计900 mm，与内部有效长度一致；C区略长以容纳风机和多气体节点"))
    append_styled_row(dimensions_table, ("接水盘", "约280 × 180 × 25 mm，可拆", "位于B区独立防水范围；50 mL演示水量下保留充足防溢余量"))
    append_styled_row(dimensions_table, ("电气舱", "可用空间不小于280 × 180 × 70 mm", "布置在A/C区底座干区；220V适配器始终位于模型外部"))
    append_styled_row(dimensions_table, ("加工控制", "总装±5 mm；安装孔±0.5 mm", "先制作1:1纸板样机；透明板和3D打印件预留1-2 mm装配间隙"))

    layout_shape = document.inline_shapes[2]
    relation_id = layout_shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    document.part.related_parts[relation_id]._blob = make_layout_diagram()

    document.core_properties.title = "综合管廊数字孪生运维实体样品项目计划书 V2.8"
    document.core_properties.subject = "1 m实体样品尺寸冻结、数字孪生与运维闭环"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
